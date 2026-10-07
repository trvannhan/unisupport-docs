# -*- coding: utf-8 -*-
"""
Script tạo file Word tổng hợp toàn bộ tài liệu dự án UniSupport.
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

def set_cell_shading(cell, color_hex):
    """Set background color for a table cell."""
    shading_elm = cell._element.get_or_add_tcPr()
    shading = shading_elm.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear',
    })
    shading_elm.append(shading)

def add_table_from_md(doc, headers, rows, col_widths=None):
    """Add a formatted table to the document."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, '2E4057')
    
    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            if r_idx % 2 == 0:
                set_cell_shading(cell, 'F0F4F8')
    
    return table

def add_section_title(doc, text, level=1):
    """Add a heading with custom formatting."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x5E)
    return heading

def read_md_file(filepath):
    """Read markdown file and return content."""
    with open(filepath, 'r', encoding='utf-8') as f:
        return f.read()

def parse_md_table(text):
    """Parse a markdown table into headers and rows."""
    lines = text.strip().split('\n')
    table_lines = [l for l in lines if '|' in l]
    if len(table_lines) < 3:
        return None, None
    
    headers = [c.strip() for c in table_lines[0].split('|') if c.strip()]
    rows = []
    for line in table_lines[2:]:  # skip header and separator
        cols = [c.strip() for c in line.split('|') if c.strip() or line.count('|') > 1]
        # Better parsing
        parts = line.split('|')
        cols = [p.strip() for p in parts[1:-1]] if len(parts) > 2 else [p.strip() for p in parts if p.strip()]
        if cols:
            rows.append(cols)
    return headers, rows

def add_md_content(doc, content, base_heading_level=2):
    """Add markdown content to doc with basic parsing."""
    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r')
        
        # Skip empty lines
        if not line.strip():
            i += 1
            continue
        
        # Headings
        if line.startswith('######'):
            add_section_title(doc, line.lstrip('#').strip(), level=min(base_heading_level + 5, 9))
            i += 1
            continue
        elif line.startswith('#####'):
            add_section_title(doc, line.lstrip('#').strip(), level=min(base_heading_level + 4, 9))
            i += 1
            continue
        elif line.startswith('####'):
            add_section_title(doc, line.lstrip('#').strip(), level=min(base_heading_level + 3, 9))
            i += 1
            continue
        elif line.startswith('###'):
            add_section_title(doc, line.lstrip('#').strip(), level=min(base_heading_level + 2, 9))
            i += 1
            continue
        elif line.startswith('##'):
            add_section_title(doc, line.lstrip('#').strip(), level=min(base_heading_level + 1, 9))
            i += 1
            continue
        elif line.startswith('#'):
            add_section_title(doc, line.lstrip('#').strip(), level=base_heading_level)
            i += 1
            continue
        
        # Table detection
        if '|' in line and i + 1 < len(lines) and '---' in lines[i+1]:
            table_lines = []
            j = i
            while j < len(lines) and '|' in lines[j].rstrip('\r'):
                table_lines.append(lines[j].rstrip('\r'))
                j += 1
            
            if len(table_lines) >= 3:
                headers_line = table_lines[0]
                headers = [c.strip() for c in headers_line.split('|')[1:-1]]
                if not headers:
                    headers = [c.strip() for c in headers_line.split('|') if c.strip()]
                
                rows = []
                for tl in table_lines[2:]:
                    parts = tl.split('|')
                    if len(parts) > 2:
                        cols = [p.strip() for p in parts[1:-1]]
                    else:
                        cols = [p.strip() for p in parts if p.strip()]
                    if cols and any(c for c in cols):
                        # Pad or trim to match headers length
                        while len(cols) < len(headers):
                            cols.append('')
                        cols = cols[:len(headers)]
                        rows.append(cols)
                
                if headers and rows:
                    add_table_from_md(doc, headers, rows)
                    doc.add_paragraph('')  # spacing
                
                i = j
                continue
        
        # Code blocks
        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].rstrip('\r').strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\r'))
                i += 1
            i += 1  # skip closing ```
            
            if code_lines:
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                p.paragraph_format.left_indent = Cm(1)
            continue
        
        # Horizontal rule
        if line.strip() == '---':
            doc.add_paragraph('─' * 60)
            i += 1
            continue
        
        # Block quote
        if line.strip().startswith('>'):
            text = line.strip().lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(text)
            run.font.size = Pt(9)
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue
        
        # Bullet/list items
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            # Clean up markdown bold/italic
            text = text.replace('**', '').replace('__', '')
            indent_level = len(line) - len(line.lstrip())
            p = doc.add_paragraph(text, style='List Bullet')
            for run in p.runs:
                run.font.size = Pt(9)
            i += 1
            continue
        
        # Numbered list
        if re.match(r'^\s*\d+\.\s', line):
            text = re.sub(r'^\s*\d+\.\s', '', line)
            text = text.replace('**', '').replace('__', '')
            p = doc.add_paragraph(text, style='List Number')
            for run in p.runs:
                run.font.size = Pt(9)
            i += 1
            continue
        
        # Regular paragraph
        text = line.strip().replace('**', '').replace('__', '').replace('`', '')
        if text:
            p = doc.add_paragraph(text)
            for run in p.runs:
                run.font.size = Pt(10)
        i += 1

def main():
    doc = Document()
    
    # Document properties
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    
    # ========================
    # TITLE PAGE
    # ========================
    for _ in range(6):
        doc.add_paragraph('')
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('TÀI LIỆU DỰ ÁN UNISUPPORT')
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x5E)
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Hệ thống Quản lý & Hỗ trợ Sinh viên\nAurora University')
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x4A, 0x6F, 0xA5)
    
    doc.add_paragraph('')
    
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run('Tổng kinh phí: 300.000.000 VNĐ\nThời gian thực hiện: 14 tuần\nQuy mô: 3.000 sinh viên')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    
    doc.add_page_break()
    
    # ========================
    # TABLE OF CONTENTS (Manual)
    # ========================
    add_section_title(doc, 'MỤC LỤC', level=1)
    
    toc_items = [
        ('Phần 1', 'CÂY CẤU TRÚC DỰ ÁN'),
        ('Phần 2', 'TỔNG QUAN SẢN PHẨM (01-product)'),
        ('Phần 3', 'NGHIỆP VỤ HỖ TRỢ SINH VIÊN (02-domain)'),
        ('Phần 4', 'ĐẶC TẢ CÁC PHÂN HỆ (03-modules)'),
        ('Phần 5', 'QUY TRÌNH NGHIỆP VỤ (04-workflows)'),
        ('Phần 6', 'YÊU CẦU PHI CHỨC NĂNG (05-nfr)'),
        ('Phần 7', 'TIÊU CHÍ NGHIỆM THU (06-acceptance)'),
        ('Phần 8', 'KIẾN TRÚC HỆ THỐNG (07-architecture)'),
        ('Phần 9', 'QUẢN LÝ DỰ ÁN (08-project)'),
    ]
    
    for num, title_text in toc_items:
        p = doc.add_paragraph()
        run = p.add_run(f'{num}: {title_text}')
        run.font.size = Pt(11)
        run.bold = True
    
    doc.add_page_break()
    
    # ========================
    # PHẦN 1: CÂY CẤU TRÚC DỰ ÁN
    # ========================
    add_section_title(doc, 'PHẦN 1: CÂY CẤU TRÚC DỰ ÁN (Documentation Structure)', level=1)
    
    tree_text = """docs/
├── README.md                          # Tổng quan toàn bộ hệ thống tài liệu
├── 01-product/                        # Tổng quan sản phẩm, mục tiêu & phạm vi
│   ├── product-overview.md            # Tổng quan dự án UniSupport
│   ├── problem-statement.md           # Hiện trạng & Khó khăn
│   ├── goals-and-non-goals.md         # Mục tiêu & Scope/Out of Scope
│   ├── actors-and-roles.md            # Vai trò (Sinh viên, Nhân viên, Quản lý)
│   └── product-scope.md               # Phạm vi sản phẩm & Các giả định
├── 02-domain/                         # Nghiệp vụ hỗ trợ sinh viên
│   ├── domain-overview.md             # Tổng quan nghiệp vụ
│   ├── terminology.md                 # Thuật ngữ nghiệp vụ
│   ├── ticket-model.md                # Mô hình dữ liệu Ticket
│   ├── ticket-lifecycle.md            # Vòng đời phiếu hỗ trợ
│   ├── state-transition.md            # Ma trận chuyển đổi trạng thái
│   └── business-rules.md              # Quy tắc nghiệp vụ
├── 03-modules/                        # Đặc tả yêu cầu chi tiết
│   ├── M01-student-portal/            # PRD Phân hệ Sinh viên
│   │   ├── README.md
│   │   └── prd.md
│   ├── M02-staff-operations/          # PRD Phân hệ Nhân viên
│   │   ├── README.md
│   │   └── prd.md
│   ├── M03-management-dashboard/      # PRD Phân hệ Quản lý/Admin
│   │   ├── README.md
│   │   └── prd.md
│   ├── M04-notification-service/      # Đặc tả thông báo nội bộ
│   │   ├── README.md
│   │   └── prd.md
│   └── M05-rbac-security/             # Đặc tả phân quyền & bảo mật
│       ├── README.md
│       └── prd.md
├── 04-workflows/                      # Quy trình nghiệp vụ
│   ├── README.md
│   ├── WF-01-submit-support-request.md
│   ├── WF-02-claim-and-triage.md
│   ├── WF-03-request-supplement.md
│   ├── WF-04-transfer-department.md
│   ├── WF-05-complete-and-resolve.md
│   └── WF-06-close-and-rate.md
├── 05-non-functional-requirements/    # Yêu cầu phi chức năng
│   ├── README.md
│   ├── performance.md
│   ├── security.md
│   ├── reliability.md
│   ├── usability.md
│   ├── observability.md
│   └── deployment.md
├── 06-acceptance/                     # Tiêu chí nghiệm thu
│   ├── README.md
│   ├── traceability-matrix.md
│   └── test-scenarios.md
├── 07-architecture/                   # Kiến trúc hệ thống
│   ├── README.md
│   ├── system-context.md
│   ├── architecture-overview.md
│   ├── data-model.md
│   ├── api-design.md
│   ├── security-design.md
│   └── technical-decisions/
│       ├── README.md
│       ├── ADR-001-ticket-id-generation.md
│       ├── ADR-002-role-based-access-control.md
│       └── ADR-003-file-attachment-storage.md
└── 08-project/                        # Quản lý dự án
    ├── assumptions.md
    ├── constraints.md
    ├── milestones-and-timeline.md
    ├── open-questions.md
    ├── glossary.md
    ├── resource-and-cost.md
    └── team-charter.md"""
    
    p = doc.add_paragraph()
    run = p.add_run(tree_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x5E)
    
    doc.add_page_break()
    
    # ========================
    # DEFINE FILE STRUCTURE TO PROCESS
    # ========================
    sections = [
        {
            'title': 'PHẦN 2: TỔNG QUAN SẢN PHẨM (01-product)',
            'files': [
                ('2.1 Tổng Quan Sản Phẩm', 'd:/docs/01-product/product-overview.md'),
                ('2.2 Hiện Trạng & Bài Toán Nghiệp Vụ', 'd:/docs/01-product/problem-statement.md'),
                ('2.3 Mục Tiêu & Phạm Vi', 'd:/docs/01-product/goals-and-non-goals.md'),
                ('2.4 Các Vai Trò & Phân Quyền', 'd:/docs/01-product/actors-and-roles.md'),
                ('2.5 Phạm Vi Sản Phẩm & Giả Định', 'd:/docs/01-product/product-scope.md'),
            ]
        },
        {
            'title': 'PHẦN 3: NGHIỆP VỤ HỖ TRỢ SINH VIÊN (02-domain)',
            'files': [
                ('3.1 Tổng Quan Nghiệp Vụ', 'd:/docs/02-domain/domain-overview.md'),
                ('3.2 Thuật Ngữ Nghiệp Vụ', 'd:/docs/02-domain/terminology.md'),
                ('3.3 Mô Hình Dữ Liệu Ticket', 'd:/docs/02-domain/ticket-model.md'),
                ('3.4 Vòng Đời Phiếu Hỗ Trợ', 'd:/docs/02-domain/ticket-lifecycle.md'),
                ('3.5 Ma Trận Chuyển Đổi Trạng Thái', 'd:/docs/02-domain/state-transition.md'),
                ('3.6 Quy Tắc Nghiệp Vụ', 'd:/docs/02-domain/business-rules.md'),
            ]
        },
        {
            'title': 'PHẦN 4: ĐẶC TẢ CÁC PHÂN HỆ (03-modules)',
            'files': [
                ('4.1 Phân Hệ Sinh Viên (M01) - Tổng Quan', 'd:/docs/03-modules/M01-student-portal/README.md'),
                ('4.2 Phân Hệ Sinh Viên (M01) - PRD Chi Tiết', 'd:/docs/03-modules/M01-student-portal/prd.md'),
                ('4.3 Phân Hệ Nhân Viên (M02) - Tổng Quan', 'd:/docs/03-modules/M02-staff-operations/README.md'),
                ('4.4 Phân Hệ Nhân Viên (M02) - PRD Chi Tiết', 'd:/docs/03-modules/M02-staff-operations/prd.md'),
                ('4.5 Phân Hệ Quản Lý (M03) - Tổng Quan', 'd:/docs/03-modules/M03-management-dashboard/README.md'),
                ('4.6 Phân Hệ Quản Lý (M03) - PRD Chi Tiết', 'd:/docs/03-modules/M03-management-dashboard/prd.md'),
                ('4.7 Thông Báo Nội Bộ (M04) - Tổng Quan', 'd:/docs/03-modules/M04-notification-service/README.md'),
                ('4.8 Thông Báo Nội Bộ (M04) - PRD Chi Tiết', 'd:/docs/03-modules/M04-notification-service/prd.md'),
                ('4.9 Phân Quyền & Bảo Mật (M05) - Tổng Quan', 'd:/docs/03-modules/M05-rbac-security/README.md'),
                ('4.10 Phân Quyền & Bảo Mật (M05) - PRD Chi Tiết', 'd:/docs/03-modules/M05-rbac-security/prd.md'),
            ]
        },
        {
            'title': 'PHẦN 5: QUY TRÌNH NGHIỆP VỤ (04-workflows)',
            'files': [
                ('5.0 Tổng Quan Quy Trình', 'd:/docs/04-workflows/README.md'),
                ('5.1 WF-01: Gửi Yêu Cầu Hỗ Trợ', 'd:/docs/04-workflows/WF-01-submit-support-request.md'),
                ('5.2 WF-02: Tiếp Nhận & Phân Loại', 'd:/docs/04-workflows/WF-02-claim-and-triage.md'),
                ('5.3 WF-03: Yêu Cầu Bổ Sung Hồ Sơ', 'd:/docs/04-workflows/WF-03-request-supplement.md'),
                ('5.4 WF-04: Chuyển Tiếp Phòng Ban', 'd:/docs/04-workflows/WF-04-transfer-department.md'),
                ('5.5 WF-05: Cập Nhật Kết Quả & Đóng', 'd:/docs/04-workflows/WF-05-complete-and-resolve.md'),
                ('5.6 WF-06: Đánh Giá Hài Lòng', 'd:/docs/04-workflows/WF-06-close-and-rate.md'),
            ]
        },
        {
            'title': 'PHẦN 6: YÊU CẦU PHI CHỨC NĂNG (05-nfr)',
            'files': [
                ('6.0 Tổng Quan NFR', 'd:/docs/05-non-functional-requirements/README.md'),
                ('6.1 Hiệu Năng', 'd:/docs/05-non-functional-requirements/performance.md'),
                ('6.2 Bảo Mật', 'd:/docs/05-non-functional-requirements/security.md'),
                ('6.3 Độ Ổn Định & Tin Cậy', 'd:/docs/05-non-functional-requirements/reliability.md'),
                ('6.4 Tính Dễ Sử Dụng', 'd:/docs/05-non-functional-requirements/usability.md'),
                ('6.5 Giám Sát & Log', 'd:/docs/05-non-functional-requirements/observability.md'),
                ('6.6 Triển Khai & Hạ Tầng', 'd:/docs/05-non-functional-requirements/deployment.md'),
            ]
        },
        {
            'title': 'PHẦN 7: TIÊU CHÍ NGHIỆM THU (06-acceptance)',
            'files': [
                ('7.0 Tổng Quan Nghiệm Thu', 'd:/docs/06-acceptance/README.md'),
                ('7.1 Ma Trận Truy Xuất Yêu Cầu (RTM)', 'd:/docs/06-acceptance/traceability-matrix.md'),
                ('7.2 Kịch Bản Kiểm Thử UAT', 'd:/docs/06-acceptance/test-scenarios.md'),
            ]
        },
        {
            'title': 'PHẦN 8: KIẾN TRÚC HỆ THỐNG (07-architecture)',
            'files': [
                ('8.0 Tổng Quan Kiến Trúc', 'd:/docs/07-architecture/README.md'),
                ('8.1 Sơ Đồ Ngữ Cảnh (System Context)', 'd:/docs/07-architecture/system-context.md'),
                ('8.2 Kiến Trúc Tổng Quan', 'd:/docs/07-architecture/architecture-overview.md'),
                ('8.3 Mô Hình Dữ Liệu & ERD', 'd:/docs/07-architecture/data-model.md'),
                ('8.4 Thiết Kế RESTful API', 'd:/docs/07-architecture/api-design.md'),
                ('8.5 Thiết Kế Bảo Mật & RBAC', 'd:/docs/07-architecture/security-design.md'),
                ('8.6 Quyết Định Kiến Trúc (ADR) - Tổng Quan', 'd:/docs/07-architecture/technical-decisions/README.md'),
                ('8.7 ADR-001: Sinh Mã Ticket', 'd:/docs/07-architecture/technical-decisions/ADR-001-ticket-id-generation.md'),
                ('8.8 ADR-002: Phân Quyền RBAC', 'd:/docs/07-architecture/technical-decisions/ADR-002-role-based-access-control.md'),
                ('8.9 ADR-003: Lưu Trữ File Đính Kèm', 'd:/docs/07-architecture/technical-decisions/ADR-003-file-attachment-storage.md'),
            ]
        },
        {
            'title': 'PHẦN 9: QUẢN LÝ DỰ ÁN (08-project)',
            'files': [
                ('9.1 Giả Định Dự Án', 'd:/docs/08-project/assumptions.md'),
                ('9.2 Ràng Buộc Dự Án', 'd:/docs/08-project/constraints.md'),
                ('9.3 Tiến Độ & Cột Mốc', 'd:/docs/08-project/milestones-and-timeline.md'),
                ('9.4 Vấn Đề Chờ Thảo Luận', 'd:/docs/08-project/open-questions.md'),
                ('9.5 Thuật Ngữ Dự Án', 'd:/docs/08-project/glossary.md'),
                ('9.6 Nguồn Lực & Chi Phí', 'd:/docs/08-project/resource-and-cost.md'),
                ('9.7 Điều Lệ Nhóm (Team Charter)', 'd:/docs/08-project/team-charter.md'),
            ]
        },
    ]
    
    # ========================
    # PROCESS EACH SECTION
    # ========================
    for section in sections:
        add_section_title(doc, section['title'], level=1)
        
        for sub_title, filepath in section['files']:
            add_section_title(doc, sub_title, level=2)
            
            try:
                content = read_md_file(filepath)
                add_md_content(doc, content, base_heading_level=3)
            except Exception as e:
                doc.add_paragraph(f'[Lỗi đọc file: {filepath}] - {str(e)}')
            
            doc.add_page_break()
    
    # ========================
    # SAVE FILE
    # ========================
    output_path = 'd:/docs/UniSupport_TaiLieuDuAn_DayDu.docx'
    doc.save(output_path)
    print(f'Da tao thanh cong file Word: {output_path}')
    print(f'Tong so phan: {sum(len(s["files"]) for s in sections)} tai lieu')

if __name__ == '__main__':
    main()
